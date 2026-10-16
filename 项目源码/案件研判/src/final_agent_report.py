from __future__ import annotations

from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from typing import Any
import json
import re
import time

from .llm_utils import cached_chat_json, model_candidates, update_audit_from_result
from .execution_policy import FallbackPolicy, normalize_fallback_policy
from .evidence_chain import is_incomplete_tool_output


STATUS_CN = {
    "enhanced": "证据增强",
    "weakened": "证据削弱",
    "unresolved": "待进一步核查",
}

TOOL_CN = {
    "entity_model_slice": "主体定向切片核查",
    "relation_neighbor_subgraph": "关系邻域核查",
    "fund_flow_path_probe": "资金路径核查",
    "same_amount_window_probe": "同日同额窗口核查",
    "counterparty_overlap_probe": "对手方重合核查",
    "transfer_structure_analyzer": "转账结构深度核查",
    "llm_judge": "大模型证据判断",
    "stop": "停止研判",
}


@dataclass(frozen=True)
class ReportAssertion:
    assertion_id: str
    model_class: str
    model_ids: list[int]
    entity_name: str | None
    title: str
    statement: str
    status: str
    severity: str
    fact_refs: list[str]
    supplement_fact_refs: list[str]
    evidence_chain_id: str
    limitations: str


def generate_final_agent_report(
    config: Any,
    case_id: int,
    model_factpacks: list[dict[str, Any]],
    model_class_packs: list[dict[str, Any]],
    case_pack: dict[str, Any],
    planner_result: dict[str, Any],
    fallback_policy: str | FallbackPolicy = FallbackPolicy.REVIEW_SAFE,
) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    policy = normalize_fallback_policy(fallback_policy)
    context = _build_final_writer_context(
        case_id,
        model_factpacks,
        model_class_packs,
        case_pack,
        planner_result,
    )
    audit = {
        "mode": "structured-final-writer",
        "status": "not_started",
        "llm_requested": bool(config.llm.is_complete),
        "call_count": 0,
        "fallback_used": False,
        "context_chars": len(json.dumps(context, ensure_ascii=False, default=str)),
        "validation_errors": [],
        "fallback_policy": policy.value,
    }
    assertions: list[ReportAssertion] = []
    if config.llm.is_complete:
        candidates = model_candidates(config.llm.model, "final_writer")[:2]
        audit["model_candidates"] = candidates
        repair_errors: list[str] = []
        for attempt in range(1, 3):
            model_name = candidates[min(attempt - 1, len(candidates) - 1)] if candidates else config.llm.model
            started = time.perf_counter()
            audit["call_count"] += 1
            try:
                prompt = _final_writer_prompt(repair_errors)
                result = cached_chat_json(
                    config,
                    model_name=model_name,
                    role=f"final-writer-{attempt}",
                    messages=[
                        {"role": "system", "content": prompt},
                        {"role": "user", "content": json.dumps(context, ensure_ascii=False, default=str)},
                    ],
                    max_tokens=min(config.llm.max_tokens, 4096),
                )
                update_audit_from_result(audit, result, model_name, attempt)
                audit["last_latency_seconds"] = round(time.perf_counter() - started, 3)
                assertions, repair_errors = _validate_assertions(result.payload, context)
                if assertions:
                    audit["status"] = "success"
                    audit["used_model"] = model_name
                    break
            except Exception as exc:
                repair_errors = [f"{type(exc).__name__}: {str(exc)[:200]}"]
                audit["last_latency_seconds"] = round(time.perf_counter() - started, 3)
        audit["validation_errors"] = repair_errors

    if not assertions:
        deterministic = _deterministic_assertions(context)
        assertions, fallback_errors = _validate_assertions(
            {"assertions": [asdict(item) for item in deterministic]},
            context,
        )
        audit["status"] = "fallback"
        audit["fallback_used"] = True
        audit["validation_errors"] = [*audit.get("validation_errors", []), *fallback_errors]
        if not audit["validation_errors"]:
            audit["validation_errors"] = ["Final Writer did not return valid assertions"]

    assertion_dicts = [asdict(item) for item in assertions]
    report = render_final_agent_report(
        case_id,
        model_factpacks,
        model_class_packs,
        case_pack,
        planner_result,
        assertion_dicts,
    )
    audit["assertion_count"] = len(assertion_dicts)
    audit["report_chars"] = len(report)
    return report, assertion_dicts, audit


def _final_writer_prompt(repair_errors: list[str]) -> str:
    repair = ""
    if repair_errors:
        repair = "上一次输出存在以下问题，请修复后重新输出完整 JSON：" + "；".join(repair_errors[:8])
    return (
        "你是案件模型研判报告的结构化断言生成器。只能使用输入中的 hypothesis、final_status、"
        "fact_refs、supplement_fact_refs 和 evidence_chain_id，不得创造主体、模型、金额、时间或关系。"
        "所有 title、statement、severity、limitations 必须使用中文。"
        "enhanced 表示证据增强，weakened 表示证据削弱，unresolved 表示待进一步核查。"
        "每条 evidence chain 输出一条 assertion，status 必须与 chain 的 final_status 完全一致。"
        "model_ids 必须完整复制当前 evidence chain 的 model_ids，不得遗漏或增加模型。"
        "只输出 JSON：{\"assertions\":[{\"assertion_id\":\"assert_001\","
        "\"model_class\":\"多维分析\",\"model_ids\":[70],\"entity_name\":\"...\","
        "\"title\":\"...\",\"statement\":\"...\",\"status\":\"enhanced|weakened|unresolved\","
        "\"severity\":\"重点关注|一般关注|待核查\",\"fact_refs\":[\"m70_entity_01\"],"
        "\"supplement_fact_refs\":[\"sf_001\"],\"evidence_chain_id\":\"chain_...\","
        "\"limitations\":\"...\"}]}。" + repair
    )


def _build_final_writer_context(
    case_id: int,
    model_factpacks: list[dict[str, Any]],
    model_class_packs: list[dict[str, Any]],
    case_pack: dict[str, Any],
    planner_result: dict[str, Any],
) -> dict[str, Any]:
    base_fact_index = _base_fact_index(model_factpacks)
    supplements_by_hyp: dict[str, list[dict[str, Any]]] = {}
    for fact in planner_result.get("supplement_facts", []):
        supplements_by_hyp.setdefault(str(fact.get("hypothesis_id")), []).append(fact)
    class_by_model: dict[int, str] = {}
    for pack in model_class_packs:
        for model in pack.get("models", []):
            if model.get("model_id") is not None:
                class_by_model[int(model["model_id"])] = str(pack.get("class_name") or "未分类")

    chains = []
    for chain in planner_result.get("evidence_chains", []):
        hypothesis = chain.get("hypothesis", {})
        hyp_id = str(hypothesis.get("hypothesis_id"))
        model_ids = [int(value) for value in hypothesis.get("target_model_ids", []) if str(value).isdigit()]
        for fact_id in hypothesis.get("basis_fact_ids", []):
            match = re.match(r"m(\d+)_", str(fact_id))
            if match:
                model_ids.append(int(match.group(1)))
        model_ids = list(dict.fromkeys(model_ids))
        supplements = supplements_by_hyp.get(hyp_id, [])
        chain_payload = {
                "evidence_chain_id": chain.get("chain_id"),
                "model_class": class_by_model.get(model_ids[0], "未分类") if model_ids else "未分类",
                "model_ids": model_ids,
                "hypothesis": {
                    "hypothesis_id": hyp_id,
                    "entity_name": hypothesis.get("entity_name"),
                    "statement": hypothesis.get("statement"),
                    "basis_fact_ids": hypothesis.get("basis_fact_ids", []),
                },
                "final_status": chain.get("final_status"),
                "supplement_facts": [
                    {
                        "supplement_fact_id": fact.get("supplement_fact_id"),
                        "tool_name": fact.get("tool_name"),
                        "summary": fact.get("output", {}).get("summary"),
                        "status_signal": fact.get("output", {}).get("status_signal"),
                        "anomaly_flags": fact.get("output", {}).get("anomaly_flags", [])[:8],
                        "fact_refs": fact.get("fact_refs", [])[:8],
                        "rows_read": fact.get("output", {}).get("rows_read"),
                        "truncated": bool(fact.get("output", {}).get("truncated")),
                        "limit_reason": fact.get("output", {}).get("limit_reason"),
                        "aggregation_complete": fact.get("output", {}).get("aggregation_complete", True),
                        "evidence_values": _claim_evidence_output(fact.get("output", {})),
                    }
                    for fact in supplements
                ],
                "judge_assessment": _last_judge_assessment(chain),
                "base_facts": [
                    base_fact_index[fact_id]
                    for fact_id in hypothesis.get("basis_fact_ids", [])
                    if fact_id in base_fact_index
                ],
            }
        chain_payload["allowed_numeric_tokens"] = sorted(_allowed_numeric_tokens(chain_payload))
        chains.append(chain_payload)
    return {
        "case_id": case_id,
        "data_source": case_pack.get("data_source"),
        "total_result_count": case_pack.get("total_result_count"),
        "known_model_ids": [int(fp.get("model_context", {}).get("model_id")) for fp in model_factpacks],
        "known_fact_ids": sorted(_known_base_fact_ids(model_factpacks)),
        "known_supplement_fact_ids": sorted(
            str(fact.get("supplement_fact_id"))
            for fact in planner_result.get("supplement_facts", [])
            if fact.get("supplement_fact_id")
        ),
        "known_entities": sorted(_known_entities(model_factpacks, chains)),
        "chains": chains,
    }


def _validate_assertions(
    payload: dict[str, Any] | None,
    context: dict[str, Any],
) -> tuple[list[ReportAssertion], list[str]]:
    if not isinstance(payload, dict) or not isinstance(payload.get("assertions"), list):
        return [], ["输出缺少 assertions 数组"]
    known_facts = set(context.get("known_fact_ids", []))
    known_supplements = set(context.get("known_supplement_fact_ids", []))
    known_models = set(context.get("known_model_ids", []))
    known_entities = set(context.get("known_entities", []))
    chains = {str(item.get("evidence_chain_id")): item for item in context.get("chains", [])}
    valid: list[ReportAssertion] = []
    errors: list[str] = []
    for index, raw in enumerate(payload.get("assertions", []), start=1):
        if not isinstance(raw, dict):
            errors.append(f"assertion {index} 不是对象")
            continue
        chain_id = str(raw.get("evidence_chain_id") or "")
        chain = chains.get(chain_id)
        status = str(raw.get("status") or "")
        entity_name = str(raw.get("entity_name") or "") or None
        model_ids = [int(value) for value in raw.get("model_ids", []) if str(value).isdigit()]
        fact_refs = [str(value) for value in raw.get("fact_refs", [])]
        supplement_refs = [str(value) for value in raw.get("supplement_fact_refs", [])]
        title = str(raw.get("title") or "").strip()
        statement = str(raw.get("statement") or "").strip()
        severity = str(raw.get("severity") or "").strip()
        limitations = str(raw.get("limitations") or "").strip()
        item_errors = []
        if chain is None:
            item_errors.append("evidence_chain_id 不存在")
        elif status != chain.get("final_status"):
            item_errors.append("status 与 evidence chain 不一致")
        if chain is not None:
            chain_fact_ids = set(chain.get("hypothesis", {}).get("basis_fact_ids", []))
            chain_supplement_ids = {
                str(item.get("supplement_fact_id"))
                for item in chain.get("supplement_facts", [])
                if item.get("supplement_fact_id")
            }
            if any(value not in chain_fact_ids for value in fact_refs):
                item_errors.append("fact_refs 不属于该证据链假设")
            if any(value not in chain_supplement_ids for value in supplement_refs):
                item_errors.append("supplement_fact_refs 不属于该证据链")
            chain_model_ids = {int(value) for value in chain.get("model_ids", [])}
            if set(model_ids) != chain_model_ids:
                item_errors.append("model_ids 必须完整匹配该证据链")
            chain_entity = str(chain.get("hypothesis", {}).get("entity_name") or "") or None
            if entity_name != chain_entity:
                item_errors.append("entity_name 与证据链主体不一致")
            if str(raw.get("model_class") or "") != str(chain.get("model_class") or ""):
                item_errors.append("model_class 与证据链不一致")
            chain_incomplete = any(is_incomplete_tool_output(item) for item in chain.get("supplement_facts", []))
            if chain_incomplete and status != "unresolved":
                item_errors.append("不完整证据链只能生成待核查断言")
            if chain_incomplete and not any(term in limitations for term in ("截断", "不完整", "未完成")):
                item_errors.append("不完整证据链的限制说明必须披露数据不完整")
        if status not in STATUS_CN:
            item_errors.append("status 非法")
        if any(value not in known_facts for value in fact_refs):
            item_errors.append("fact_refs 包含未知引用")
        if any(value not in known_supplements for value in supplement_refs):
            item_errors.append("supplement_fact_refs 包含未知引用")
        if not fact_refs and not supplement_refs:
            item_errors.append("缺少事实引用")
        if any(value not in known_models for value in model_ids) or not model_ids:
            item_errors.append("model_ids 非法")
        if entity_name and entity_name not in known_entities:
            item_errors.append("entity_name 不在已知实体中")
        if not all(_contains_chinese(value) for value in [title, statement, severity, limitations]):
            item_errors.append("中文字段为空或不含中文")
        statement_numbers = _text_numeric_tokens(title + " " + statement + " " + limitations)
        chain_numbers = set(chain.get("allowed_numeric_tokens", [])) if chain is not None else set()
        if not statement_numbers.issubset(chain_numbers):
            item_errors.append("断言包含当前证据链中不存在的数字")
        if chain is not None:
            chain_entity = str(chain.get("hypothesis", {}).get("entity_name") or "")
            claim_text = " ".join((title, statement, limitations))
            foreign_entities = {
                name for name in known_entities if len(name) >= 2 and name != chain_entity and name in claim_text
            }
            if foreign_entities:
                item_errors.append("断言包含其他证据链主体")
        if item_errors:
            errors.append(f"assertion {index}: " + "；".join(item_errors))
            continue
        valid.append(
            ReportAssertion(
                assertion_id=str(raw.get("assertion_id") or f"assert_{index:03d}"),
                model_class=str(raw.get("model_class") or chain.get("model_class") or "未分类"),
                model_ids=list(dict.fromkeys(model_ids)),
                entity_name=entity_name,
                title=title,
                statement=statement,
                status=status,
                severity=severity,
                fact_refs=fact_refs,
                supplement_fact_refs=supplement_refs,
                evidence_chain_id=chain_id,
                limitations=limitations,
            )
        )
    if len(valid) != len(chains):
        errors.append(f"有效断言数量 {len(valid)} 与证据链数量 {len(chains)} 不一致")
        return [], errors
    return valid, errors


def _deterministic_assertions(context: dict[str, Any]) -> list[ReportAssertion]:
    result = []
    for index, chain in enumerate(context.get("chains", []), start=1):
        hypothesis = chain.get("hypothesis", {})
        supplements = chain.get("supplement_facts", [])
        status = str(chain.get("final_status") or "unresolved")
        incomplete = any(is_incomplete_tool_output(item) for item in supplements)
        if incomplete:
            status = "unresolved"
        entity_name = hypothesis.get("entity_name")
        result.append(
            ReportAssertion(
                assertion_id=f"assert_{index:03d}",
                model_class=str(chain.get("model_class") or "未分类"),
                model_ids=[int(value) for value in chain.get("model_ids", [])],
                entity_name=str(entity_name) if entity_name else None,
                title=f"{entity_name or '相关主体'}研判线索",
                statement=(
                    f"当前补证结果增强了对{entity_name or '相关主体'}相关线索的支持。"
                    if status == "enhanced"
                    else f"关于{entity_name or '相关主体'}的线索仍需结合后续证据继续核查。"
                    if status == "unresolved"
                    else f"当前补证结果削弱了对{entity_name or '相关主体'}相关线索的支持。"
                ),
                status=status if status in STATUS_CN else "unresolved",
                severity="重点关注" if status == "enhanced" else "待核查",
                fact_refs=[str(value) for value in hypothesis.get("basis_fact_ids", [])],
                supplement_fact_refs=[
                    str(item.get("supplement_fact_id"))
                    for item in supplements
                    if item.get("supplement_fact_id")
                ],
                evidence_chain_id=str(chain.get("evidence_chain_id") or ""),
                limitations=(
                    "补证工具结果达到预算上限并发生截断，当前仅保留为待核查线索。"
                    if incomplete
                    else "该断言为结构化事实与证据链的保守表达，不超出当前可用证据范围。"
                ),
            )
        )
    return result


def render_final_agent_report(
    case_id: int,
    model_factpacks: list[dict[str, Any]],
    model_class_packs: list[dict[str, Any]],
    case_pack: dict[str, Any],
    planner_result: dict[str, Any],
    assertions: list[dict[str, Any]],
) -> str:
    by_status = {status: [item for item in assertions if item.get("status") == status] for status in STATUS_CN}
    incomplete_facts = [
        fact
        for fact in planner_result.get("supplement_facts", [])
        if is_incomplete_tool_output(fact.get("output"))
    ]
    model_names = _model_name_map(model_factpacks)
    clean_assertions = [_normalize_assertion_text(item, model_names) for item in assertions]
    by_status = {status: [item for item in clean_assertions if item.get("status") == status] for status in STATUS_CN}
    enhanced_model_ids = {
        int(model_id)
        for item in by_status["enhanced"]
        for model_id in item.get("model_ids", [])
        if str(model_id).isdigit()
    }
    case_info = case_pack.get("case_info") or {}
    event_name = case_info.get("event_name") or f"案例 {case_id}"
    event_code = case_info.get("event_code") or "未提供"
    event_phase = _case_phase_text(case_info.get("event_phase"))
    total_count = int(case_pack.get("total_result_count") or 0)
    lines = [
        f"# 案例 {case_id} 模型研判报告",
        "",
        "> 本报告以模型结果为主线。基础数字由程序确定性提取，研判线索由受控多轮研判流程补证并经证据判断。",
        "",
        "## 一、案例概况",
        "",
        f"事件名称为 **{event_name}**，事件编号为 **{event_code}**，当前处于{event_phase}。"
        "本次分析围绕相关自然人及其银行账户活动展开，重点观察账户余额、现金存取、资金流向、"
        "转账关系以及同日同额交易等信息；报告中列出的主体仅为本次重点关注对象，并非案件涉及人员的完整清单。",
        "",
        f"本报告选取 {len(model_factpacks)} 个代表模型，共覆盖 **{_fmt_int(total_count)} 条**模型结果。"
        f"当前形成 {len(clean_assertions)} 条可追溯研判线索，其中证据增强 {len(by_status['enhanced'])} 条、"
        f"待进一步核查 {len(by_status['unresolved'])} 条、证据削弱 {len(by_status['weakened'])} 条。"
        "报告以模型事实为依据，对跨模型重复出现且获得进一步证据支持的线索进行归纳，供后续核查使用，"
        "不直接将模型信号认定为违法违规事实。",
        "",
        "## 二、分析范围与总体画像",
        "",
        "| 模型 | 业务用途 | 结果规模 | 报告角色 |",
        "|---|---|---:|---|",
    ]
    for fp in model_factpacks:
        ctx = fp.get("model_context", {})
        model_id = int(ctx.get("model_id") or 0)
        result_count = int(ctx.get("result_count") or 0)
        lines.append(
            f"| {model_id} {_clean_model_name(ctx.get('model_name'))} | "
            f"{_model_business_use(ctx, fp)} | {_fmt_int(result_count)} | "
            f"{_report_role(model_id, result_count, enhanced_model_ids)} |"
        )

    lines.extend(["", _overall_portrait(model_factpacks, enhanced_model_ids), ""])

    lines.extend(["## 三、核心研判发现", ""])
    if by_status["enhanced"]:
        for index, item in enumerate(by_status["enhanced"], start=1):
            refs = [*item.get("fact_refs", []), *item.get("supplement_fact_refs", [])]
            lines.append(f"### {index}. {item.get('title')}")
            lines.append("")
            lines.append(f"**{STATUS_CN.get(str(item.get('status')), item.get('status'))}**")
            lines.append("")
            lines.append(_trim_terminal_punctuation(str(item.get("statement") or "")) + "。")
            lines.append("")
            lines.append(
                "证据导航："
                + "；".join([f"基础事实 {ref}" for ref in item.get("fact_refs", [])])
                + (("；" if refs and item.get("supplement_fact_refs") else "") + "；".join(f"补证 {ref}" for ref in item.get("supplement_fact_refs", [])))
                + f"；证据链 {item.get('evidence_chain_id')}。"
            )
            lines.append("")
    else:
        lines.append("- 当前没有证据增强状态的核心研判发现。")
        lines.append("")

    lines.extend(["## 四、重点模型分析", ""])
    for fp in _primary_factpacks(model_factpacks, enhanced_model_ids):
        _append_model_analysis(lines, fp, model_names, clean_assertions)

    lines.extend(["## 五、风险归纳与核查建议", ""])
    _append_risk_summary(lines, by_status, model_names, planner_result)

    lines.extend(["## 六、证据索引与附录", ""])
    lines.extend(["### 6.1 完整证据链", ""])
    for chain in planner_result.get("evidence_chains", []):
        hypothesis = chain.get("hypothesis", {})
        lines.append(
            f"#### 证据链 {chain.get('chain_id')}：{hypothesis.get('entity_name') or '未指定主体'}"
        )
        lines.append(f"- 待核查线索：{_normalize_model_mentions(str(hypothesis.get('statement') or ''), model_names)}")
        lines.append(f"- 最终状态：{STATUS_CN.get(str(chain.get('final_status')), chain.get('final_status'))}")
        for step in chain.get("steps", []):
            decision = step.get("decision", {})
            tool = str(decision.get("tool_name") or "")
            lines.append(
                f"- 第 {step.get('round')} 步：{TOOL_CN.get(tool, tool)}；"
                f"状态：{STATUS_CN.get(str(step.get('status')), step.get('status'))}；"
                f"{_normalize_model_mentions(_localize_internal_terms(str(step.get('assessment') or '')), model_names)}"
            )
        lines.append("")

    lines.extend(["### 6.2 基础事实索引", ""])
    for fp in model_factpacks:
        for entity in fp.get("focus_entities", [])[:5]:
            _append_base_fact(lines, fp, entity, "entity")
        for relation in fp.get("relation_profiles", [])[:3]:
            _append_base_fact(lines, fp, relation, "relation")

    lines.extend(["### 6.3 补证结果索引", ""])
    for fact in planner_result.get("supplement_facts", []):
        output = fact.get("output", {})
        lines.append(f"#### 补证 {fact.get('supplement_fact_id')}")
        lines.append(f"- 工具：{TOOL_CN.get(str(fact.get('tool_name')), fact.get('tool_name'))}。")
        lines.append(f"- 范围：{_normalize_model_mentions(str(output.get('summary') or '未提供摘要'), model_names)}")
        lines.append(f"- 结果状态：{STATUS_CN.get(str(output.get('status_signal')), output.get('status_signal'))}。")
        if is_incomplete_tool_output(output):
            lines.append("- 限制说明：该补证结果存在截断或聚合未完成，仅可作为待进一步核查信息。")
        lines.append("")

    lines.extend(
        [
            "### 6.4 数据限制与方法说明",
            "",
            f"- 数据来源：{'真实数据库' if case_pack.get('data_source') == 'database' else case_pack.get('data_source')}；"
            "本次统计基于已缓存的模型结果。",
            "- 大模型不直接访问数据库、不生成 SQL，只能在压缩事实和白名单工具范围内决策。",
            "- 补证事实与证据链采用追加记录；基础模型事实包不被大模型修改。",
            "- 未获得充分补证的线索不会写成确定性异常结论。",
            f"- 本次共有 {len(incomplete_facts)} 个补证结果被截断或聚合未完成；不完整结果仅保留为待核查信息，不形成增强或削弱结论。",
            "",
        ]
    )
    return "\n".join(lines)


def _model_name_map(model_factpacks: list[dict[str, Any]]) -> dict[int, str]:
    return {
        int(fp.get("model_context", {}).get("model_id")): _clean_model_name(fp.get("model_context", {}).get("model_name"))
        for fp in model_factpacks
        if fp.get("model_context", {}).get("model_id") is not None
    }


def _case_phase_text(value: Any) -> str:
    raw = str(value or "").strip()
    return {
        "1": "分析事件信息阶段",
        "2": "获取分析数据阶段",
        "3": "构建分析模型阶段",
        "4": "研判阶段",
        "5": "制作分析报告阶段",
    }.get(raw, raw or "研判阶段")


def _clean_model_name(value: Any) -> str:
    return str(value or "").lstrip("★").strip() or "未命名模型"


def _normalize_assertion_text(item: dict[str, Any], model_names: dict[int, str]) -> dict[str, Any]:
    result = dict(item)
    for key in ("title", "statement", "limitations"):
        result[key] = _normalize_model_mentions(str(result.get(key) or ""), model_names)
    return result


def _normalize_model_mentions(text: str, model_names: dict[int, str]) -> str:
    result = text
    for model_id, model_name in sorted(model_names.items(), key=lambda item: len(str(item[0])), reverse=True):
        result = re.sub(
            rf"模型\s*{model_id}(?!\s*[）\)])",
            f"{model_name}（模型 {model_id}）",
            result,
        )
    return result


def _fmt_int(value: Any) -> str:
    try:
        return f"{int(value):,}"
    except Exception:
        return str(value)


def _fmt_amount(value: Any) -> str:
    try:
        number = Decimal(str(value))
    except Exception:
        return str(value)
    if number == number.to_integral():
        return f"{int(number):,}"
    return f"{number:,.2f}"


def _report_role(model_id: int, result_count: int, enhanced_model_ids: set[int]) -> str:
    if result_count <= 0:
        return "当前无结果"
    if model_id in enhanced_model_ids:
        return "重点证据模型"
    return "覆盖概览"


def _model_business_use(ctx: dict[str, Any], factpack: dict[str, Any]) -> str:
    desc = str(ctx.get("demo_role") or ctx.get("model_desc") or "").strip()
    if desc:
        return desc
    if factpack.get("relation_profiles"):
        return "观察关系、金额、频次和对手方结构"
    if factpack.get("focus_entities"):
        return "观察主体、账户、金额和时间分布"
    return "用于范围描述和索引"


def _overall_portrait(model_factpacks: list[dict[str, Any]], enhanced_model_ids: set[int]) -> str:
    largest = sorted(
        [
            (
                int(fp.get("model_context", {}).get("result_count") or 0),
                int(fp.get("model_context", {}).get("model_id") or 0),
                _clean_model_name(fp.get("model_context", {}).get("model_name")),
            )
            for fp in model_factpacks
        ],
        reverse=True,
    )
    primary = [
        f"{_clean_model_name(fp.get('model_context', {}).get('model_name'))}（模型 {fp.get('model_context', {}).get('model_id')}）"
        for fp in model_factpacks
        if int(fp.get("model_context", {}).get("model_id") or 0) in enhanced_model_ids
    ]
    if not largest:
        return "当前没有可汇总的模型结果。"
    top = largest[0]
    primary_text = "、".join(primary) if primary else "当前没有证据增强状态的重点证据模型"
    return (
        f"从结果规模看，{top[2]}（模型 {top[1]}）构成主要数据来源，共 {_fmt_int(top[0])} 条结果；"
        f"{primary_text}承担本次核心研判发现的主要证据支撑。其余模型用于范围描述、基础索引或补充观察。"
    )


def _primary_factpacks(model_factpacks: list[dict[str, Any]], enhanced_model_ids: set[int]) -> list[dict[str, Any]]:
    primary = [
        fp
        for fp in model_factpacks
        if int(fp.get("model_context", {}).get("model_id") or 0) in enhanced_model_ids
    ]
    return primary or [
        fp for fp in model_factpacks if int(fp.get("model_context", {}).get("result_count") or 0) > 0
    ][:5]


def _append_model_analysis(
    lines: list[str],
    factpack: dict[str, Any],
    model_names: dict[int, str],
    assertions: list[dict[str, Any]],
) -> None:
    ctx = factpack.get("model_context", {})
    model_id = int(ctx.get("model_id") or 0)
    model_name = _clean_model_name(ctx.get("model_name"))
    related = [item for item in assertions if model_id in [int(v) for v in item.get("model_ids", []) if str(v).isdigit()]]
    focus_name = related[0].get("entity_name") if related else None
    lines.append(f"### {model_name}（模型 {model_id}）")
    lines.append(f"- **模型用途：** {_model_business_use(ctx, factpack)}。")
    lines.append(f"- **结果规模：** {_fmt_int(ctx.get('result_count', 0))} 条。")
    if focus_name:
        lines.append(f"- **关注对象：** {focus_name}。")
    summary = _factpack_key_summary(factpack, focus_name)
    if summary:
        lines.append(f"- **关键事实：** {summary}")
    if related:
        lines.append(f"- **业务提示：** {_model_specific_business_tip(model_id, model_name, summary, related[0], model_names)}")
        refs = [*related[0].get("fact_refs", []), *related[0].get("supplement_fact_refs", [])]
        lines.append(f"- **证据导航：** {'；'.join(refs)}；证据链 {related[0].get('evidence_chain_id')}。")
    lines.append("")


def _model_specific_business_tip(
    model_id: int,
    model_name: str,
    summary: str,
    assertion: dict[str, Any],
    model_names: dict[int, str],
) -> str:
    model_ids = [int(mid) for mid in assertion.get("model_ids", []) if str(mid).isdigit()]
    others = [mid for mid in model_ids if mid != model_id]
    title = _trim_terminal_punctuation(str(assertion.get("title") or "当前研判线索"))
    base = _trim_terminal_punctuation(summary) if summary else f"{model_name}提供本模型下的基础事实"
    if others:
        other_text = "、".join(f"{model_names.get(mid, f'模型 {mid}')}（模型 {mid}）" for mid in others)
        return f"{model_name}（模型 {model_id}）中，{base}；该事实与{other_text}共同支撑“{title}”。"
    return f"{model_name}（模型 {model_id}）中，{base}；该事实为“{title}”提供模型内重点支撑。"


def _factpack_key_summary(factpack: dict[str, Any], focus_name: Any) -> str:
    candidates = []
    if focus_name:
        candidates.extend(
            entity for entity in factpack.get("focus_entities", []) if entity.get("entity_name") == focus_name
        )
    candidates.extend(factpack.get("focus_entities", [])[:1])
    if candidates:
        entity = candidates[0]
        metrics = entity.get("metrics", {})
        parts = [f"{entity.get('entity_name')}相关基础事实包含 {_fmt_int(metrics.get('record_count', 0))} 条记录"]
        if metrics.get("amount_sum") is not None:
            parts.append(f"金额合计 {_fmt_amount(metrics.get('amount_sum'))} 元")
        if metrics.get("amount_max") is not None:
            parts.append(f"单笔最高 {_fmt_amount(metrics.get('amount_max'))} 元")
        if metrics.get("counterparty_count") is not None:
            parts.append(f"涉及 {_fmt_int(metrics.get('counterparty_count'))} 个对手方")
        return "，".join(parts) + "。"
    relations = factpack.get("relation_profiles", [])
    if relations:
        relation = relations[0]
        return (
            f"重点关系 {relation.get('source')} → {relation.get('target')}，"
            f"记录 {_fmt_int(relation.get('count', 0))} 条，金额合计 {_fmt_amount(relation.get('amount_sum', 0))} 元。"
        )
    return ""


def _append_risk_summary(
    lines: list[str],
    by_status: dict[str, list[dict[str, Any]]],
    model_names: dict[int, str],
    planner_result: dict[str, Any],
) -> None:
    supplements = _supplements_by_id(planner_result)
    lines.append("### 5.1 风险归纳")
    if not by_status["enhanced"]:
        lines.append("- 当前没有证据增强状态的线索，暂不形成重点风险归纳。")
    for item in by_status["enhanced"]:
        model_ids = [int(mid) for mid in item.get("model_ids", []) if str(mid).isdigit()]
        model_text = "、".join(
            f"{model_names.get(int(mid), f'模型 {mid}')}（模型 {int(mid)}）"
            for mid in model_ids
        )
        support_text = "形成交叉支撑" if len(model_ids) >= 2 else "提供重点支撑"
        lines.append(
            f"- **{item.get('entity_name') or '相关主体'}：** {model_text}{support_text}，"
            f"当前状态为{STATUS_CN.get(str(item.get('status')), item.get('status'))}；"
            f"建议作为后续核查优先线索。"
        )
    if by_status["unresolved"]:
        lines.append(f"- 另有 {len(by_status['unresolved'])} 条线索仍处于待进一步核查状态，不进入确定性风险归纳。")
    if by_status["weakened"]:
        lines.append(f"- 另有 {len(by_status['weakened'])} 条线索被当前补证结果削弱或排除，应降低核查优先级。")
    lines.append("")
    lines.append("### 5.2 核查建议")
    lines.append("- 调取重点账户的完整原始流水，核对交易时间、金额、方向、摘要和余额变化，确认模型结果与原始记录一致。")
    for item in by_status["enhanced"]:
        entity = item.get("entity_name") or "相关主体"
        model_text = "、".join(
            model_names.get(int(mid), f"模型 {mid}")
            for mid in item.get("model_ids", [])
            if str(mid).isdigit()
        )
        detail = _assertion_supplement_detail(item, supplements, model_names)
        suffix = f"重点关注{detail}。" if detail else "核验主要对手方、账户归属、资金来源去向和真实交易背景。"
        lines.append(f"- 围绕{entity}复核{model_text}中的重点记录，{suffix}")
    lines.append("- 对余额、现金存取、转账关系等模型信号保持审慎表达，避免仅凭模型命中结果直接作出性质判断。")
    lines.append("")


def _supplements_by_id(planner_result: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(item.get("supplement_fact_id")): item
        for item in planner_result.get("supplement_facts", [])
        if item.get("supplement_fact_id")
    }


def _assertion_supplement_detail(
    assertion: dict[str, Any],
    supplements: dict[str, dict[str, Any]],
    model_names: dict[int, str],
) -> str:
    highlights = []
    for supplement_id in assertion.get("supplement_fact_refs", []):
        supplement = supplements.get(str(supplement_id))
        if not supplement:
            continue
        highlight = _supplement_highlight(supplement.get("output", {}), model_names)
        if highlight:
            highlights.append(highlight)
    return "；".join(highlights[:2])


def _supplement_highlight(output: dict[str, Any], model_names: dict[int, str]) -> str:
    if output.get("top_outgoing") or output.get("top_incoming"):
        parts = []
        for label, key in (("主要流出", "top_outgoing"), ("主要流入", "top_incoming")):
            rows = output.get(key) or []
            if rows:
                top = rows[0]
                parts.append(
                    f"{label}{top.get('source')}→{top.get('target')} {_fmt_int(top.get('count', 0))} 条、"
                    f"{_fmt_amount(top.get('amount_sum', 0))} 元"
                )
        return "，".join(parts)
    entities = output.get("focus_entities") or []
    if entities:
        entity = entities[0]
        metrics = entity.get("metrics", {})
        parts = [
            f"{entity.get('entity_name')} {_fmt_int(metrics.get('record_count', 0))} 条",
        ]
        if metrics.get("amount_sum") is not None:
            parts.append(f"金额合计 {_fmt_amount(metrics.get('amount_sum'))} 元")
        if metrics.get("amount_max") is not None:
            parts.append(f"单笔最高 {_fmt_amount(metrics.get('amount_max'))} 元")
        time_text = _time_bucket_text(entity.get("top_time_buckets", {}))
        if time_text:
            parts.append(time_text)
        return "、".join(parts)
    summary = _normalize_model_mentions(str(output.get("summary") or ""), model_names)
    return _trim_terminal_punctuation(summary)


def _time_bucket_text(buckets: dict[str, Any]) -> str:
    hours = buckets.get("hours") or []
    if not hours:
        return ""
    top = hours[0]
    return f"高峰时段 {top.get('value')} 时段 {_fmt_int(top.get('count', 0))} 条"


def _append_base_fact(lines: list[str], factpack: dict[str, Any], item: dict[str, Any], kind: str) -> None:
    ctx = factpack.get("model_context", {})
    fact_id = item.get("fact_id")
    if not fact_id:
        return
    model_name = _clean_model_name(ctx.get("model_name"))
    lines.append(f"#### 基础事实 {fact_id}")
    lines.append(f"- 模型：{model_name}（模型 {ctx.get('model_id')}）。")
    if kind == "entity":
        metrics = item.get("metrics", {})
        lines.append(
            f"- 对象：{item.get('entity_name')}；记录 {_fmt_int(metrics.get('record_count', 0))} 条"
            + (f"；金额合计 {_fmt_amount(metrics.get('amount_sum'))} 元" if metrics.get("amount_sum") is not None else "")
            + "。"
        )
    else:
        lines.append(
            f"- 关系：{item.get('source')} → {item.get('target')}；记录 {_fmt_int(item.get('count', 0))} 条"
            + (f"；金额合计 {_fmt_amount(item.get('amount_sum'))} 元" if item.get("amount_sum") is not None else "")
            + "。"
        )
    evidence_records = item.get("evidence_records") or [
        record.get("evidence") for record in item.get("top_amount_records", []) if record.get("evidence")
    ]
    if evidence_records:
        lines.append(f"- 代表性记录：{'; '.join(_compact_record(record) for record in evidence_records[:3])}。")
    lines.append("")


def _compact_record(record: Any) -> str:
    if not isinstance(record, dict):
        return str(record)
    parts = []
    keys = (
        "source_row_id",
        "row_id",
        "id",
        "姓名",
        "本方户名",
        "流出对象（起点）",
        "流入对象（终点）",
        "source",
        "target",
        "卡号/账号",
        "卡号",
        "本方银行卡号",
        "流出账户",
        "流入账户",
        "对手账户",
        "对端银行卡号清单",
        "交易对手",
        "对端户名清单",
        "amount",
        "交易金额",
        "金额",
        "转出金额",
        "个人银行账户总余额(元)",
        "最后一次交易后的账户余额(元)",
        "方向",
        "交易时间",
        "交易日期",
        "time",
        "摘要",
        "同日相同金额转不同对端数量",
    )
    for key in keys:
        if record.get(key) not in [None, ""]:
            parts.append(f"{key}={_short_value(record.get(key))}")
    return "，".join(parts[:5]) or str(record)[:120]


def _short_value(value: Any, max_chars: int = 36) -> str:
    text = str(value)
    return text if len(text) <= max_chars else text[: max_chars - 1] + "…"


def write_final_report_docx(markdown: str, path: Any) -> None:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size, color in [
        ("Heading 1", 16, "0B2545"),
        ("Heading 2", 13, "1F4D78"),
        ("Heading 3", 11.5, "2E74B5"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    def set_cell_shading(cell: Any, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def add_table(block: list[str]) -> None:
        headers = [cell.strip() for cell in block[0].strip("|").split("|")]
        rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in block[2:]]
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = _plain_docx_text(header)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F2F4F7")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row[: len(headers)]):
                cells[index].text = _plain_docx_text(value)
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cells[index].paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
        doc.add_paragraph()

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        if line.startswith("<a "):
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index])
                index += 1
            if len(block) >= 2:
                add_table(block)
            continue
        if line.startswith("# "):
            doc.add_heading(_plain_docx_text(line[2:]), level=0)
        elif line.startswith("## "):
            doc.add_heading(_plain_docx_text(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(_plain_docx_text(line[4:]), level=2)
        elif line.startswith("#### "):
            doc.add_heading(_plain_docx_text(line[5:]), level=3)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(_plain_docx_text(line[2:]))
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(_plain_docx_text(line[2:]))
            run.italic = True
            run.font.color.rgb = RGBColor.from_string("475467")
        else:
            doc.add_paragraph(_plain_docx_text(line))
        index += 1
    doc.save(path)


def _plain_docx_text(text: str) -> str:
    result = re.sub(r"`([^`]+)`", r"\1", text or "")
    result = re.sub(r"\*\*([^*]+)\*\*", r"\1", result)
    result = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", result)
    result = re.sub(r"<[^>]+>", "", result)
    return result


def _known_base_fact_ids(model_factpacks: list[dict[str, Any]]) -> set[str]:
    result: set[str] = set()
    for fp in model_factpacks:
        for key in ("focus_entities", "relation_profiles"):
            for item in fp.get(key, []):
                if item.get("fact_id"):
                    result.add(str(item["fact_id"]))
    return result


def _base_fact_index(model_factpacks: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for factpack in model_factpacks:
        model_id = factpack.get("model_context", {}).get("model_id")
        for key in ("focus_entities", "relation_profiles"):
            for item in factpack.get(key, []):
                fact_id = str(item.get("fact_id") or "")
                if fact_id:
                    result[fact_id] = {
                        "fact_id": fact_id,
                        "model_id": model_id,
                        **_claim_evidence_output(item),
                    }
    return result


def _claim_evidence_output(output: dict[str, Any]) -> dict[str, Any]:
    keys = (
        "summary",
        "row_count",
        "numeric_summary",
        "metrics",
        "entity_name",
        "source",
        "target",
        "count",
        "amount_sum",
        "top_outgoing",
        "top_incoming",
        "top_same_amount_groups",
        "direct_outgoing_top",
        "two_hop_paths",
        "overlaps",
        "focus_entities",
        "relation_profiles",
        "structured",
        "anomaly_flags",
    )
    result: dict[str, Any] = {}
    for key in keys:
        value = output.get(key)
        if isinstance(value, list):
            result[key] = value[:5]
        elif value not in [None, "", [], {}]:
            result[key] = value
    return result


def _allowed_numeric_tokens(value: Any, key_path: str = "") -> set[str]:
    result: set[str] = set()
    if isinstance(value, dict):
        for key, item in value.items():
            lowered = str(key).lower()
            if any(marker in lowered for marker in ("_id", "_ids", "_ref", "_refs", "fingerprint")):
                if lowered == "model_ids":
                    result.update(_allowed_numeric_tokens(item, lowered))
                continue
            result.update(_allowed_numeric_tokens(item, f"{key_path}.{lowered}"))
        return result
    if isinstance(value, (list, tuple)):
        for item in value:
            result.update(_allowed_numeric_tokens(item, key_path))
        return result
    if isinstance(value, bool) or value is None:
        return result
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
        except InvalidOperation:
            return result
        result.update(_decimal_forms(number))
        if any(marker in key_path for marker in ("amount", "金额", "资金")):
            result.update(_scaled_decimal_forms(number, Decimal("10000")))
            result.update(_scaled_decimal_forms(number, Decimal("100000000")))
        return result
    if isinstance(value, str) and any(
        marker in key_path
        for marker in ("summary", "statement", "assessment", "limitations", "date", "time")
    ):
        result.update(_text_numeric_tokens(value))
    return result


def _text_numeric_tokens(text: str) -> set[str]:
    result = set()
    for raw in re.findall(r"\d[\d,]*(?:\.\d+)?", text or ""):
        try:
            result.add(_decimal_token(Decimal(raw.replace(",", ""))))
        except InvalidOperation:
            continue
    return result


def _decimal_forms(number: Decimal) -> set[str]:
    return {_decimal_token(number), *(_decimal_token(number.quantize(Decimal("1").scaleb(-p))) for p in range(3))}


def _scaled_decimal_forms(number: Decimal, divisor: Decimal) -> set[str]:
    scaled = number / divisor
    return {_decimal_token(scaled.quantize(Decimal("1").scaleb(-p))) for p in range(3)}


def _decimal_token(number: Decimal) -> str:
    if number == number.to_integral():
        return str(number.quantize(Decimal("1")))
    return format(number.normalize(), "f")


def _known_entities(model_factpacks: list[dict[str, Any]], chains: list[dict[str, Any]]) -> set[str]:
    result: set[str] = set()
    for fp in model_factpacks:
        for entity in fp.get("focus_entities", []):
            if entity.get("entity_name"):
                result.add(str(entity["entity_name"]))
        for relation in fp.get("relation_profiles", []):
            for key in ("source", "target"):
                if relation.get(key):
                    result.add(str(relation[key]))
    for chain in chains:
        entity = chain.get("hypothesis", {}).get("entity_name")
        if entity:
            result.add(str(entity))
    return result


def _last_judge_assessment(chain: dict[str, Any]) -> str:
    for step in reversed(chain.get("steps", [])):
        if step.get("decision", {}).get("tool_name") == "llm_judge":
            return str(step.get("assessment") or "")
    return ""


def _contains_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def _localize_internal_terms(text: str) -> str:
    replacements = {
        "top outgoing": "主要流出关系",
        "top incoming": "主要流入关系",
        "top_outgoing": "主要流出关系",
        "top_incoming": "主要流入关系",
        "row_count": "记录条数",
        "truncated": "结果截断",
        "aggregation_complete": "聚合完整",
        "enhanced": "证据增强",
        "weakened": "证据削弱",
        "unresolved": "待进一步核查",
    }
    result = text
    for source, target in replacements.items():
        result = re.sub(re.escape(source), target, result, flags=re.I)
    return result


def _trim_terminal_punctuation(text: str) -> str:
    return text.rstrip("。；; ")
