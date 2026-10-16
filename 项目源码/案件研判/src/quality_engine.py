"""End-to-end offline engine and reproducible artifact writer."""

from __future__ import annotations

import hashlib
import html
import json
import os
import tempfile
import time
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from .quality_models import Claim, Fact, ModelProfile, Observation, RunManifest, TraceEvent, utc_now
from .quality_operators import compute_facts
from .quality_llm import BoundedLLM
from .quality_cross_model import build_cross_model_links
from .quality_context import pack_context
from .quality_reasoning import ToolRegistry, build_claims, run_policy, validate_claims
from .quality_registry import DEFAULT_DATA_DIR, build_registry, validate_registry


def _json_default(value: Any) -> Any:
    if hasattr(value, "to_dict"):
        return value.to_dict()
    raise TypeError(f"Unsupported JSON value: {type(value)!r}")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, payload: Any) -> None:
    # Reuse the unique-temp text writer so JSON has the same cleanup and
    # collision guarantees as checkpoints and projections.
    _write_text_atomic(
        path,
        json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default),
    )


def _write_text_atomic(path: Path, content: str) -> None:
    """Write a text artifact without exposing a partially rewritten file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            handle.write(content)
        temporary.replace(path)
    except Exception:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def _write_docx_atomic(report: dict[str, Any], path: Path) -> None:
    """Generate a DOCX off-path, then publish the closed ZIP atomically."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        _docx(report, temporary)
        temporary.replace(path)
    except Exception:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def _allocate_run_dir(output_root: Path) -> tuple[str, Path]:
    """Reserve a collision-safe run directory before writing any artifacts.

    The human-readable timestamp remains the primary run id.  A bounded
    numeric suffix is added only when another process has already reserved the
    same second, and ``exist_ok=False`` makes the reservation race-safe.
    """
    output_root.mkdir(parents=True, exist_ok=True)
    base = f"quality_{time.strftime('%Y%m%d_%H%M%S')}"
    for suffix in range(1000):
        run_id = base if suffix == 0 else f"{base}_{suffix:03d}"
        run_dir = output_root / run_id
        try:
            run_dir.mkdir(parents=True, exist_ok=False)
        except FileExistsError:
            continue
        return run_id, run_dir
    raise RuntimeError(f"could not allocate a unique run directory under {output_root}")


def _checkpoint_payload(
    *,
    status: str,
    run_id: str,
    manifest: RunManifest,
    model_items: list[dict[str, Any]],
    processed_model_ids: set[int],
    facts: list[Fact],
    observations: list[Any],
    claims: list[Claim],
    trace: list[TraceEvent],
    elapsed_seconds: float,
    resume_count: int = 0,
    last_model_id: int | None = None,
    include_state: bool = True,
    state_file: str = "checkpoint_models.jsonl",
    selected_model_ids: set[int] | None = None,
) -> dict[str, Any]:
    """Serialize enough in-flight state to continue after an interrupted run.

    The checkpoint deliberately stores typed pipeline state rather than only a
    progress counter.  A process can therefore restart without recomputing
    completed models or silently dropping their evidence chain.
    """
    payload = {
        "status": status,
        "run_id": run_id,
        "last_model_id": last_model_id,
        "processed_models": len(processed_model_ids),
        "processed_model_ids": sorted(processed_model_ids),
        "updated_at": utc_now(),
        "resume_count": resume_count,
        "elapsed_seconds": elapsed_seconds,
        "manifest": manifest.to_dict(),
        "state_file": state_file,
        "selected_model_ids": sorted(selected_model_ids) if selected_model_ids is not None else None,
    }
    if include_state:
        payload.update(
            {
                "model_items": model_items,
                "facts": [fact.to_dict() for fact in facts],
                "observations": [observation.to_dict() for observation in observations],
                "claims": [claim.to_dict() for claim in claims],
                "trace": [event.to_dict() for event in trace],
            }
        )
    return payload


def _markdown(report: dict[str, Any]) -> str:
    manifest = report["manifest"]
    coverage = report["coverage"]
    models = report.get("models", []) if isinstance(report.get("models", []), list) else []
    facts = report.get("facts", []) if isinstance(report.get("facts", []), list) else []
    observations = report.get("observations", []) if isinstance(report.get("observations", []), list) else []
    claims = report.get("claims", []) if isinstance(report.get("claims", []), list) else []
    trace = report.get("trace", []) if isinstance(report.get("trace", []), list) else []

    def cell(value: Any) -> str:
        """Keep dynamic values safe and readable inside Markdown tables."""
        return str(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ")

    fact_by_id = {
        fact.get("fact_id"): fact
        for fact in facts
        if isinstance(fact, dict) and isinstance(fact.get("fact_id"), str)
    }
    sampled_models = [
        item for item in models
        if isinstance(item, dict) and item.get("coverage") == "sampled"
    ]
    unsupported_models = [
        item for item in models
        if isinstance(item, dict) and item.get("coverage") == "unsupported_no_sample"
    ]
    policy_counts: dict[str, int] = {}
    trace_counts: dict[str, int] = {}
    for event in trace:
        if not isinstance(event, dict):
            continue
        event_type = str(event.get("event_type") or "unknown")
        trace_counts[event_type] = trace_counts.get(event_type, 0) + 1
        if event_type == "policy_selected":
            payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
            policy = str(payload.get("policy") or "unknown")
            policy_counts[policy] = policy_counts.get(policy, 0) + 1
    observation_counts: dict[str, int] = {}
    for observation in observations:
        if isinstance(observation, dict):
            status = str(observation.get("status") or "unknown")
            observation_counts[status] = observation_counts.get(status, 0) + 1
    unresolved_claims = sum(
        1 for claim in claims if isinstance(claim, dict) and claim.get("level") == "unresolved"
    )
    claim_validation = (
        report.get("validation", {}).get("claims", {})
        if isinstance(report.get("validation"), dict)
        and isinstance(report.get("validation", {}).get("claims"), dict)
        else {}
    )
    context_events = [
        event for event in trace
        if isinstance(event, dict) and event.get("event_type") == "context_packed"
    ]
    context_ratios: list[float] = []
    context_budget_exceeded = 0
    for event in context_events:
        payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
        try:
            context_ratios.append(float(payload.get("compression_ratio", 1.0)))
        except (TypeError, ValueError):
            pass
        if payload.get("budget_exceeded"):
            context_budget_exceeded += 1
    average_context_ratio = sum(context_ratios) / len(context_ratios) if context_ratios else None
    budget = report.get("llm_budget") if isinstance(report.get("llm_budget"), dict) else {}
    tick = chr(96)
    lines = [
        "# Model Insight Agent 证据优先研究报告",
        "",
        f"> Run: `{manifest['run_id']}`  |  Snapshot: `{manifest['snapshot_sha256'][:12]}`",
        "",
        "## 1. 执行范围",
        "",
        f"- 案件: {', '.join(str(case_id) for case_id in manifest['case_ids']) or '全量抽样快照'}",
        f"- 模型目录: {coverage['total_models']} 个；有样本 {coverage['sampled_models']} 个；无样本安全降级 {coverage['unsupported_no_sample']} 个",
        f"- 在线 LLM: {'启用' if manifest['online'] else '未启用（离线确定性回放）'}",
        "- 结论边界：事实、规则命中、推断和待核验分级展示；本报告不作法律性质判断。",
        "",
        "## 2. 研判摘要",
        "",
        f"本次运行处理 {len(sampled_models)} 个有样本模型，形成 {len(facts)} 条 Fact、{len(claims)} 条 Claim；其中待核验 Claim {unresolved_claims} 条，Observation {len(observations)} 条。",
        "",
        "| 关注项 | 当前结果 | 解释 |",
        "|---|---:|---|",
        f"| 有样本模型 | {len(sampled_models)} | 进入确定性算子和受控策略 |",
        f"| 无样本模型 | {len(unsupported_models)} | 仅展示安全降级，不生成推测结论 |",
        f"| 确定性事实 Fact | {len(facts)} | 由本地算子产生，可回到原始行 |",
        f"| Claim | {len(claims)} | 必须引用 Fact/Observation |",
        f"| 待核验 Claim | {unresolved_claims} | 不确定性不会被自动升级 |",
        "",
        "## 3. 模型覆盖",
        "",
        "| 模型 | 名称 | 覆盖状态 | 推理策略 | 事实数 | Claim数 |",
        "|---:|---|---|---|---:|---:|",
    ]
    for item in models:
        lines.append(
            f"| {cell(item.get('model_id'))} | {cell(item.get('name'))} | {cell(item.get('coverage'))} | {cell(item.get('reasoning_mode'))} | {cell(item.get('fact_count'))} | {cell(item.get('claim_count'))} |"
        )
    lines.extend(["", "## 4. 证据结论", ""])
    for claim in report["claims"]:
        marker = {"fact": "事实", "rule_hit": "规则命中", "inference": "推断", "unresolved": "待核验"}.get(claim["level"], claim["level"])
        refs = ", ".join(f"`{fact_id}`" for fact_id in claim["fact_ids"]) or "无"
        lines.append(f"- **[{marker}]** {claim['text']}（证据：{refs}）")
    if not report["claims"]:
        lines.append("- 当前没有可接受的 Claim；请查看模型覆盖和限制项。")
    lines.extend(
        [
            "",
            "## 5. 证据索引",
            "",
            "| Claim | 层级/状态 | Fact 引用 | 来源行数 | 产出方式 |",
            "|---|---|---|---:|---|",
        ]
    )
    for claim in claims:
        if not isinstance(claim, dict):
            continue
        claim_fact_ids = claim.get("fact_ids") if isinstance(claim.get("fact_ids"), list) else []
        referenced_facts = [
            fact_by_id[fact_id]
            for fact_id in claim_fact_ids
            if fact_id in fact_by_id
        ]
        source_rows = sum(
            len(fact.get("source_row_ids", []))
            for fact in referenced_facts
            if isinstance(fact.get("source_row_ids"), list)
        )
        provenance = ", ".join(
            sorted({str(fact.get("provenance") or "unknown") for fact in referenced_facts})
        ) or "无"
        refs = ", ".join(f"{tick}{fact_id}{tick}" for fact_id in claim_fact_ids) or "无"
        lines.append(
            f"| {tick}{cell(claim.get('claim_id'))}{tick} | "
            f"{cell(claim.get('level'))}/{cell(claim.get('status'))} | {refs} | "
            f"{source_rows} | {cell(provenance)} |"
        )
    referenced_fact_ids: list[str] = []
    for claim in claims:
        if not isinstance(claim, dict):
            continue
        raw_ids = claim.get("fact_ids") if isinstance(claim.get("fact_ids"), list) else []
        for fact_id in raw_ids:
            if isinstance(fact_id, str) and fact_id in fact_by_id and fact_id not in referenced_fact_ids:
                referenced_fact_ids.append(fact_id)
    lines.extend(
        [
            "",
            "## 6. 引用事实摘要",
            "",
            "| Fact | 模型/案件 | 类型/主题 | 值摘要 | 来源行（前 8 条） | 完整性 |",
            "|---|---|---|---|---|---|",
        ]
    )
    for fact_id in referenced_fact_ids:
        fact = fact_by_id[fact_id]
        raw_value = json.dumps(fact.get("value"), ensure_ascii=False, separators=(",", ":"))
        value_text = raw_value if len(raw_value) <= 220 else raw_value[:217] + "..."
        source_ids = fact.get("source_row_ids") if isinstance(fact.get("source_row_ids"), list) else []
        source_text = ",".join(str(item) for item in source_ids[:8]) or "无"
        if len(source_ids) > 8:
            source_text += f"...（另有 {len(source_ids) - 8} 条）"
        lines.append(
            f"| {tick}{cell(fact_id)}{tick} | "
            f"{cell(fact.get('model_id'))}/{cell(fact.get('case_id'))} | "
            f"{cell(fact.get('kind'))}/{cell(fact.get('subject'))} | "
            f"{cell(value_text)} | {cell(source_text)} | "
            f"{'完整' if fact.get('complete') is True else '不完整'} |"
        )
    if not referenced_fact_ids:
        lines.append("| 当前没有被 Claim 引用的 Fact | - | - | - | - | - |")
    lines.extend(["", "## 7. 策略与运行轨迹", "", "| 项目 | 计数 |", "|---|---:|"])
    for policy, count in sorted(policy_counts.items()):
        lines.append(f"| policy {tick}{cell(policy)}{tick} | {count} |")
    for status, count in sorted(observation_counts.items()):
        lines.append(f"| Observation {tick}{cell(status)}{tick} | {count} |")
    for event_type, count in sorted(trace_counts.items()):
        lines.append(f"| trace {tick}{cell(event_type)}{tick} | {count} |")
    if not policy_counts and not observation_counts and not trace_counts:
        lines.append("| 当前没有策略或轨迹事件 | 0 |")
    lines.extend(["", "## 8. 跨模型复核入口", ""])
    if report.get("cross_model_links"):
        for link in report["cross_model_links"][:30]:
            lines.append(
                f"- `{link['link_id']}`：实体线索“{link['entity_value']}”出现在模型 {', '.join(map(str, link['model_ids']))}；状态：待人工核验。"
            )
    else:
        lines.append("- 当前快照没有满足保守连接条件的跨模型候选。")
    lines.extend(["", "## 9. 限制与待核验", ""])
    if unsupported_models:
        unsupported_ids = ", ".join(
            f"{tick}{item.get('model_id')}{tick}" for item in unsupported_models
        )
        lines.append(f"- 无样本模型（安全降级）：{unsupported_ids}")
    else:
        lines.append("- 当前没有无样本模型。")
    lines.extend(
        [
            "- 抽样结果不是完整生产数据，不能据此估计真实召回率或阈值效果。",
            "- 无样本模型不会被 LLM 推测性补全。",
            "- 复杂关系结论必须回到模型原始结果和人工核验。",
            "",
            "## 10. 审计摘要",
            "",
        ]
    )
    validation = report["validation"]
    lines.extend([
        (
            f"- 上下文打包：{len(context_events)} 个事件；平均压缩比：{average_context_ratio:.6f}；预算溢出：{context_budget_exceeded} 次"
            if average_context_ratio is not None
            else "- 上下文打包：无事件"
        ),
        f"- 在线预算：{budget.get('calls_used', 0)}/{budget.get('max_calls', 0)} 次调用，{budget.get('tokens_used', 0)}/{budget.get('max_tokens', 0)} Token",
        f"- Claim 引用校验：{'通过' if validation['claims']['valid'] else '失败'}",
        f"- 运行状态：`{manifest['status']}`",
        f"- 生成耗时：{manifest.get('elapsed_seconds', 0):.3f} 秒",
    ])
    return "\n".join(lines) + "\n"


def _html(report: dict[str, Any], markdown: str) -> str:
    # The HTML is deliberately self-contained: a run directory can be copied
    # to another machine without losing the audit view.
    esc = html.escape
    manifest = report["manifest"]
    rows = "".join(
        "<tr>" + "".join(
            f"<td>{esc(str(value))}</td>" for value in (
                item["model_id"], item["name"], item["coverage"], item["reasoning_mode"], item["fact_count"], item["claim_count"]
            )
        ) + "</tr>"
        for item in report["models"]
    )
    claims = "".join(
        f"<li><span class=level>{esc(claim['level'])}</span> {esc(claim['text'])}<small>证据：{esc(', '.join(claim['fact_ids']) or '无')}</small></li>"
        for claim in report["claims"]
    ) or "<li>当前没有可接受的 Claim。</li>"
    links = "".join(
        f"<li><code>{esc(link['link_id'])}</code> 实体线索“{esc(link['entity_value'])}”出现在模型 {esc(', '.join(map(str, link['model_ids'])))}；状态：待人工核验。</li>"
        for link in report.get("cross_model_links", [])[:100]
    ) or "<li>当前没有满足保守连接条件的跨模型候选。</li>"
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Model Insight Agent</title>
<style>body{{font-family:Segoe UI,'Microsoft YaHei',sans-serif;max-width:1180px;margin:2rem auto;padding:0 1rem;color:#17202a;line-height:1.6}}h1,h2{{color:#17202a}}.badge{{display:inline-block;background:#e8f1ff;color:#164b8a;padding:.15rem .5rem;border-radius:999px}}table{{border-collapse:collapse;width:100%;margin:1rem 0;font-size:.92rem}}th,td{{border:1px solid #d9e0e8;padding:.45rem .55rem;text-align:left;vertical-align:top}}th{{background:#174a7e;color:#fff}}tr:nth-child(even){{background:#f7f9fb}}li{{margin:.45rem 0}}.level{{font-weight:600;color:#174a7e}}small{{display:block;color:#5b6775}}code{{background:#f1f3f5;padding:.08rem .25rem;border-radius:3px}}</style></head>
<body><h1>Model Insight Agent 证据优先研究报告</h1><p class="badge">{esc(manifest['status'])}</p>
<p>Run <code>{esc(manifest['run_id'])}</code> · 快照 <code>{esc(manifest['snapshot_sha256'][:12])}</code> · 案件 {esc(', '.join(map(str, manifest['case_ids'])) or '全量抽样快照')}</p>
<h2>模型覆盖</h2><table><thead><tr><th>模型</th><th>名称</th><th>状态</th><th>策略</th><th>事实</th><th>Claim</th></tr></thead><tbody>{rows}</tbody></table>
<h2>证据结论</h2><ul>{claims}</ul><h2>跨模型复核入口</h2><ul>{links}</ul>
<h2>限制与审计</h2><ul><li>抽样结果不是完整生产数据，不能据此估计真实召回率或阈值效果。</li><li>无样本模型不会被 LLM 推测性补全。</li><li>Claim 引用校验：{esc('通过' if report['validation']['claims']['valid'] else '失败')}；耗时：{manifest.get('elapsed_seconds', 0):.3f} 秒。</li></ul>
</body></html>"""


def _docx(report: dict[str, Any], path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:  # pragma: no cover - optional artifact dependency
        raise RuntimeError("python-docx is required for DOCX rendering") from exc
    document = Document()
    title = document.add_heading(level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Model Insight Agent")
    title_run.add_break()
    title_run.add_text("证据优先研究报告")
    title_run.font.size = Pt(23)
    document.add_paragraph(
        f"Run {report['manifest']['run_id']} · 快照 {report['manifest']['snapshot_sha256'][:12]}"
    )
    document.add_heading("执行范围", level=1)
    coverage = report["coverage"]
    document.add_paragraph(
        f"模型目录 {coverage['total_models']} 个；有样本 {coverage['sampled_models']} 个；无样本安全降级 {coverage['unsupported_no_sample']} 个。"
    )
    document.add_paragraph("结论按事实、规则命中、推断和待核验分级，不作法律性质判断。")
    document.add_heading("模型覆盖", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ["模型", "名称", "状态", "策略", "事实", "Claim"]):
        cell.text = text
    for item in report["models"]:
        cells = table.add_row().cells
        for cell, text in zip(cells, [
            str(item["model_id"]), item["name"], item["coverage"], item["reasoning_mode"],
            str(item["fact_count"]), str(item["claim_count"]),
        ]):
            cell.text = text
    document.add_heading("证据结论", level=1)
    for claim in report["claims"]:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(f"[{claim['level']}] {claim['text']}")
        run.font.size = Pt(10)
        refs = ", ".join(claim["fact_ids"]) or "无"
        paragraph.add_run(f" 证据：{refs}")
    document.add_heading("跨模型复核入口", level=1)
    if report.get("cross_model_links"):
        for link in report["cross_model_links"][:30]:
            document.add_paragraph(
                f"{link['link_id']}：实体线索“{link['entity_value']}”出现在模型 {', '.join(map(str, link['model_ids']))}；状态：待人工核验。",
                style="List Bullet",
            )
    else:
        document.add_paragraph("当前快照没有满足保守连接条件的跨模型候选。")
    document.add_heading("限制与审计", level=1)
    document.add_paragraph("抽样结果不代表生产全量数据；无样本模型不会被推测性补全。")
    document.add_paragraph(
        f"Claim 引用校验：{'通过' if report['validation']['claims']['valid'] else '失败'}；耗时：{report['manifest'].get('elapsed_seconds', 0):.3f} 秒。"
    )
    document.save(path)


def write_artifacts(report: dict[str, Any], output_root: Path | str) -> Path:
    output_root = Path(output_root)
    run_dir = output_root / report["manifest"]["run_id"]
    run_dir.mkdir(parents=True, exist_ok=True)
    markdown = _markdown(report)
    _write_json(run_dir / "report.json", report)
    _write_text_atomic(run_dir / "report.md", markdown)
    _write_text_atomic(run_dir / "report.html", _html(report, markdown))
    _write_docx_atomic(report, run_dir / "report.docx")
    _write_text_atomic(
        run_dir / "trace.jsonl",
        "".join(json.dumps(event, ensure_ascii=False) + "\n" for event in report["trace"]),
    )
    _write_json(run_dir / "run_manifest.json", report["manifest"])
    return run_dir


class QualityEngine:
    def __init__(
        self,
        data_dir: Path | str = DEFAULT_DATA_DIR,
        output_root: Path | str | None = None,
        online: bool = False,
        max_steps: int = 3,
    ) -> None:
        self.data_dir = Path(data_dir)
        self.output_root = Path(output_root or Path(__file__).resolve().parents[1] / "output" / "quality_runs")
        self.online = online
        self.max_steps = max_steps
        self.llm = BoundedLLM(self.output_root / ".llm_cache") if online else None

    def run(
        self,
        case_id: int | None = None,
        model_ids: set[int] | None = None,
        *,
        resume_dir: Path | str | None = None,
        stop_after_models: int | None = None,
    ) -> tuple[dict[str, Any], Path]:
        """Run the pipeline, optionally continuing a running checkpoint.

        ``stop_after_models`` is intentionally test-oriented: it lets the
        recovery contract simulate a process interruption after a durable
        checkpoint has been written.  Production callers should omit it.
        """
        started = time.perf_counter()
        profiles, rows = build_registry(self.data_dir)
        registry = validate_registry(profiles)
        sample_file = self.data_dir / "ch_model_info_data_sample.jsonl"
        resume_checkpoint_hint: dict[str, Any] | None = None
        effective_case_id = case_id
        effective_model_ids = set(model_ids) if model_ids is not None else None
        if effective_model_ids is not None:
            if not effective_model_ids:
                raise ValueError("model_ids must contain at least one model id")
            known_model_ids = {profile.model_id for profile in profiles}
            unknown_model_ids = sorted(effective_model_ids - known_model_ids)
            if unknown_model_ids:
                raise ValueError(
                    "model_ids contains unknown models: "
                    + ", ".join(str(model_id) for model_id in unknown_model_ids)
                )
        if resume_dir is not None:
            try:
                resume_checkpoint_hint = json.loads(
                    (Path(resume_dir) / "checkpoint.json").read_text(encoding="utf-8")
                )
            except (OSError, json.JSONDecodeError) as exc:
                raise ValueError(f"cannot read resume checkpoint: {Path(resume_dir) / 'checkpoint.json'}") from exc
            hint_manifest = dict(resume_checkpoint_hint.get("manifest") or {})
            hint_cases = list(hint_manifest.get("case_ids") or [])
            if effective_case_id is None and len(hint_cases) == 1:
                effective_case_id = int(hint_cases[0])
            elif effective_case_id is not None and hint_cases and sorted(int(item) for item in hint_cases) != [int(effective_case_id)]:
                raise ValueError("resume case filter does not match the running checkpoint")
            hint_models = resume_checkpoint_hint.get("selected_model_ids")
            if effective_model_ids is None and hint_models is not None:
                effective_model_ids = {int(item) for item in hint_models}
            elif hint_models is not None and effective_model_ids != {int(item) for item in hint_models}:
                raise ValueError("resume model filter does not match the running checkpoint")
            elif hint_models is None and effective_model_ids is not None:
                raise ValueError("resume model filter does not match the running checkpoint")
        grouped: dict[tuple[int, int], list[dict[str, Any]]] = defaultdict(list)
        selected_cases: set[int] = set()
        for row in rows:
            row_case = int(row["case_id"])
            if effective_case_id is not None and row_case != effective_case_id:
                continue
            grouped[(int(row["model_id"]), row_case)].append(row)
            selected_cases.add(row_case)
        if effective_case_id is not None:
            selected_cases.add(effective_case_id)

        mode_counts: dict[str, int] = defaultdict(int)
        for profile in profiles:
            mode_counts[profile.reasoning_mode] += 1
        current_snapshot_hash = _sha256(sample_file)
        resume_count = 0
        prior_elapsed = 0.0
        processed_model_ids: set[int] = set()
        state_file = "checkpoint_models.jsonl"
        if resume_dir is not None:
            run_dir = Path(resume_dir)
            checkpoint = resume_checkpoint_hint or {}
            if checkpoint.get("status") != "running":
                raise ValueError("only a running checkpoint can be resumed")
            checkpoint_manifest = dict(checkpoint.get("manifest") or {})
            if checkpoint_manifest.get("snapshot_sha256") != current_snapshot_hash:
                raise ValueError("resume checkpoint snapshot hash does not match current data")
            run_id = str(checkpoint.get("run_id") or checkpoint_manifest.get("run_id") or run_dir.name)
            manifest = RunManifest(**checkpoint_manifest)
            if manifest.online != self.online:
                raise ValueError("resume mode must match the checkpoint online flag")
            processed_model_ids = {int(item) for item in checkpoint.get("processed_model_ids", [])}
            resume_count = int(checkpoint.get("resume_count") or 0) + 1
            prior_elapsed = float(checkpoint.get("elapsed_seconds") or 0.0)
            state_file = str(checkpoint.get("state_file") or state_file)
            state_path = run_dir / state_file
            state_records: list[dict[str, Any]] = []
            if state_path.exists():
                state_text = state_path.read_text(encoding="utf-8")
                state_has_content = bool(state_text.strip())
                seen_state_model_ids: set[int] = set()
                for line in state_text.splitlines():
                    if line.strip():
                        try:
                            record = json.loads(line)
                        except json.JSONDecodeError:
                            # A process may have died while appending the last
                            # line.  The committed checkpoint is authoritative;
                            # discard the torn tail and recompute that model.
                            break
                        if not isinstance(record, dict):
                            raise ValueError("checkpoint state record must be an object")
                        try:
                            record_model_id = int(record.get("model_id"))
                        except (TypeError, ValueError) as exc:
                            raise ValueError("checkpoint state record has an invalid model_id") from exc
                        if record_model_id in processed_model_ids:
                            if record_model_id in seen_state_model_ids:
                                raise ValueError(f"checkpoint state contains duplicate committed model {record_model_id}")
                            state_records.append(record)
                            seen_state_model_ids.add(record_model_id)
                if state_has_content:
                    missing_state_model_ids = sorted(processed_model_ids - seen_state_model_ids)
                    if missing_state_model_ids:
                        raise ValueError(
                            "checkpoint state is missing committed models: "
                            + ", ".join(str(item) for item in missing_state_model_ids)
                        )
                # Remove uncommitted/torn tail records before appending new
                # deltas, preventing duplicate model IDs after a resume.
                _write_text_atomic(
                    state_path,
                    "".join(
                        json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n"
                        for record in state_records
                    ),
                )
            if state_records:
                model_items = [dict(record["model_item"]) for record in state_records]
                facts = [Fact(**item) for record in state_records for item in record.get("facts", [])]
                observations = [Observation(**item) for record in state_records for item in record.get("observations", [])]
                claims = [Claim(**item) for record in state_records for item in record.get("claims", [])]
                trace = [TraceEvent(**item) for record in state_records for item in record.get("trace", [])]
            else:
                # Backward-compatible recovery for a checkpoint created before
                # the append-only state file was introduced.
                facts = [Fact(**item) for item in checkpoint.get("facts", [])]
                observations = [Observation(**item) for item in checkpoint.get("observations", [])]
                claims = [Claim(**item) for item in checkpoint.get("claims", [])]
                trace = [TraceEvent(**item) for item in checkpoint.get("trace", [])]
                model_items = [dict(item) for item in checkpoint.get("model_items", [])]
            run_dir.mkdir(parents=True, exist_ok=True)
            if self.llm is not None:
                budget = checkpoint.get("llm_budget") or {}
                self.llm.ledger.calls_used = int(budget.get("calls_used") or manifest.llm_calls or 0)
                self.llm.ledger.tokens_used = int(budget.get("tokens_used") or manifest.llm_tokens or 0)
        else:
            run_id, run_dir = _allocate_run_dir(self.output_root)
            command = f"run --case-id {effective_case_id if effective_case_id is not None else 'all'}"
            if effective_model_ids is not None:
                command += " --model-ids " + ",".join(str(model_id) for model_id in sorted(effective_model_ids))
            manifest = RunManifest(
                run_id=run_id,
                started_at=utc_now(),
                finished_at=None,
                command=command,
                case_ids=sorted(selected_cases),
                snapshot_sha256=current_snapshot_hash,
                profile_count=len(profiles),
                sampled_profile_count=registry["sampled_models"],
                unsupported_profile_count=registry["unsupported_no_sample"],
                reasoning_modes=dict(sorted(mode_counts.items())),
                online=self.online,
            )
            facts = []
            observations = []
            claims = []
            trace = []
            model_items = []
            _write_text_atomic(run_dir / state_file, "")
            _write_json(
                run_dir / "checkpoint.json",
                _checkpoint_payload(
                    status="running",
                    run_id=run_id,
                    manifest=manifest,
                    model_items=model_items,
                    processed_model_ids=processed_model_ids,
                    facts=facts,
                    observations=observations,
                    claims=claims,
                    trace=trace,
                    elapsed_seconds=0.0,
                    resume_count=resume_count,
                    include_state=False,
                    state_file=state_file,
                    selected_model_ids=effective_model_ids,
                ),
            )
        by_model: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for (model_id, _row_case), model_rows in grouped.items():
            by_model[model_id].extend(model_rows)
        tools = ToolRegistry()
        checkpoint_interval = 5

        state_path = run_dir / state_file

        def persist_checkpoint(
            last_model_id: int | None = None,
            *,
            model_item: dict[str, Any],
            fact_start: int,
            observation_start: int,
            claim_start: int,
            trace_start: int,
        ) -> None:
            # Append one model-sized delta instead of serializing the complete
            # accumulated state on every checkpoint.  This keeps the normal
            # all-snapshot path close to its offline performance baseline.
            with state_path.open("a", encoding="utf-8") as handle:
                handle.write(
                    json.dumps(
                        {
                            "model_id": last_model_id,
                            "model_item": model_item,
                            "facts": [fact.to_dict() for fact in facts[fact_start:]],
                            "observations": [observation.to_dict() for observation in observations[observation_start:]],
                            "claims": [claim.to_dict() for claim in claims[claim_start:]],
                            "trace": [event.to_dict() for event in trace[trace_start:]],
                        },
                        ensure_ascii=False,
                        separators=(",", ":"),
                    )
                    + "\n"
                )
            force_checkpoint = stop_after_models is not None and len(processed_model_ids) >= stop_after_models
            if len(processed_model_ids) % checkpoint_interval and not force_checkpoint:
                return
            checkpoint_state = _checkpoint_payload(
                status="running",
                run_id=run_id,
                manifest=manifest,
                model_items=model_items,
                processed_model_ids=processed_model_ids,
                facts=facts,
                observations=observations,
                claims=claims,
                trace=trace,
                elapsed_seconds=prior_elapsed + (time.perf_counter() - started),
                resume_count=resume_count,
                last_model_id=last_model_id,
                include_state=False,
                state_file=state_file,
                selected_model_ids=effective_model_ids,
            )
            if self.llm is not None:
                checkpoint_state["llm_budget"] = self.llm.ledger.to_dict()
            _write_json(run_dir / "checkpoint.json", checkpoint_state)
            if stop_after_models is not None and len(processed_model_ids) >= stop_after_models:
                raise RuntimeError(f"intentional checkpoint interruption after {len(processed_model_ids)} models: {run_dir}")

        for profile in profiles:
            if profile.model_id in processed_model_ids:
                continue
            fact_start = len(facts)
            observation_start = len(observations)
            claim_start = len(claims)
            trace_start = len(trace)
            if effective_model_ids is not None and profile.model_id not in effective_model_ids:
                model_items.append({
                    "model_id": profile.model_id,
                    "name": profile.name,
                    "strategy_key": profile.strategy_key,
                    "family": profile.family,
                    "operators": list(profile.operators),
                    "specialized_tools": list(profile.specialized_tools),
                    "report_sections": list(profile.report_sections),
                    "cross_model_dependencies": list(profile.cross_model_dependencies),
                    "coverage": profile.coverage,
                    "reasoning_mode": profile.reasoning_mode,
                    "fact_count": 0,
                    "claim_count": 0,
                    "status": "not_selected",
                    "reason": "model filter",
                })
                processed_model_ids.add(profile.model_id)
                persist_checkpoint(
                    profile.model_id,
                    model_item=model_items[-1],
                    fact_start=fact_start,
                    observation_start=observation_start,
                    claim_start=claim_start,
                    trace_start=trace_start,
                )
                continue
            model_rows = by_model.get(profile.model_id, [])
            if profile.coverage == "unsupported_no_sample":
                model_items.append({
                    "model_id": profile.model_id,
                    "name": profile.name,
                    "strategy_key": profile.strategy_key,
                    "family": profile.family,
                    "operators": list(profile.operators),
                    "specialized_tools": list(profile.specialized_tools),
                    "report_sections": list(profile.report_sections),
                    "cross_model_dependencies": list(profile.cross_model_dependencies),
                    "coverage": profile.coverage,
                    "reasoning_mode": profile.reasoning_mode,
                    "fact_count": 0,
                    "claim_count": 0,
                    "status": "skipped",
                    "reason": profile.unsupported_reason,
                })
                processed_model_ids.add(profile.model_id)
                persist_checkpoint(
                    profile.model_id,
                    model_item=model_items[-1],
                    fact_start=fact_start,
                    observation_start=observation_start,
                    claim_start=claim_start,
                    trace_start=trace_start,
                )
                continue
            fact_pack = compute_facts(profile, model_rows, effective_case_id if effective_case_id is not None else "all")
            model_facts_before = len(facts)
            facts.extend(fact_pack)
            model_observations, model_trace = run_policy(
                profile,
                model_rows,
                fact_pack,
                effective_case_id if effective_case_id is not None else "all",
                tools,
                self.max_steps,
            )
            observations.extend(model_observations)
            trace.extend(model_trace)
            model_claims = build_claims(profile, fact_pack, model_observations)
            claims.extend(model_claims)
            context_pack = pack_context(profile, fact_pack, model_observations)
            trace.append(
                TraceEvent(
                    event_id=f"context_{profile.model_id}_{effective_case_id if effective_case_id is not None else 'all'}",
                    event_type="context_packed",
                    timestamp=utc_now(),
                    step=len(trace) + 1,
                    payload={
                        "model_id": profile.model_id,
                        "raw_chars": context_pack["raw_chars"],
                        "packed_chars": context_pack["packed_chars"],
                        "max_chars": context_pack["max_chars"],
                        "budget_exceeded": context_pack["budget_exceeded"],
                        "compression_ratio": context_pack["compression_ratio"],
                        "fact_ids": context_pack["fact_ids"],
                        "omitted_value_fact_ids": context_pack["omitted_value_fact_ids"],
                    },
                )
            )
            if self.llm is not None:
                advisory = self.llm.advisory_plan(profile.to_dict(), context_pack["fact_summaries"])
                trace.append(
                    TraceEvent(
                        event_id=f"llm_{profile.model_id}_{effective_case_id if effective_case_id is not None else 'all'}",
                        event_type="llm_advisory_plan",
                        timestamp=utc_now(),
                        step=len(trace) + 1,
                        payload={
                            "model_id": profile.model_id,
                            "status": advisory.status,
                            "actions": advisory.payload.get("actions", []),
                            "reason": advisory.reason,
                            "cache_key": advisory.cache_key,
                        },
                    )
                )
                manifest.llm_calls += advisory.calls_used
                manifest.llm_tokens += advisory.tokens_used
            model_items.append({
                "model_id": profile.model_id,
                "name": profile.name,
                "strategy_key": profile.strategy_key,
                "family": profile.family,
                "operators": list(profile.operators),
                "specialized_tools": list(profile.specialized_tools),
                "report_sections": list(profile.report_sections),
                "cross_model_dependencies": list(profile.cross_model_dependencies),
                "context_chars": context_pack["packed_chars"],
                "context_compression_ratio": context_pack["compression_ratio"],
                "coverage": profile.coverage,
                "reasoning_mode": profile.reasoning_mode,
                "row_count": len(model_rows),
                "fact_count": len(facts) - model_facts_before,
                "claim_count": len(model_claims),
                "status": "processed",
            })
            processed_model_ids.add(profile.model_id)
            persist_checkpoint(
                profile.model_id,
                model_item=model_items[-1],
                fact_start=fact_start,
                observation_start=observation_start,
                claim_start=claim_start,
                trace_start=trace_start,
            )
        cross_links, cross_claims = build_cross_model_links(facts)
        claims.extend(cross_claims)
        if self.llm is not None:
            # Failed/blocked calls still consume a provider attempt and must
            # remain visible in the run budget, even when no response tokens
            # were returned.
            manifest.llm_calls = self.llm.ledger.calls_used
            manifest.llm_tokens = self.llm.ledger.tokens_used
        validation = {"claims": validate_claims(claims, facts), "registry": registry}
        elapsed = prior_elapsed + (time.perf_counter() - started)
        manifest.finished_at = utc_now()
        manifest.status = "success" if validation["claims"]["valid"] else "failed"
        payload = {
            "schema_version": "1.0",
            "manifest": {**manifest.to_dict(), "elapsed_seconds": elapsed},
            "coverage": registry,
            "models": model_items,
            "facts": [fact.to_dict() for fact in facts],
            "observations": [observation.to_dict() for observation in observations],
            "claims": [claim.to_dict() for claim in claims],
            "cross_model_links": cross_links,
            "trace": [event.to_dict() for event in trace],
            "validation": validation,
            "llm_budget": self.llm.ledger.to_dict() if self.llm is not None else {
                "max_calls": 200,
                "max_tokens": 2_000_000,
                "calls_used": 0,
                "tokens_used": 0,
            },
        }
        run_dir = write_artifacts(payload, run_dir.parent)
        final_checkpoint = _checkpoint_payload(
            status=manifest.status,
            run_id=run_id,
            manifest=manifest,
            model_items=model_items,
            processed_model_ids=processed_model_ids,
            facts=facts,
            observations=observations,
            claims=claims,
            trace=trace,
            elapsed_seconds=elapsed,
            resume_count=resume_count,
            include_state=False,
            state_file=state_file,
            selected_model_ids=effective_model_ids,
        )
        final_checkpoint["completed_at"] = manifest.finished_at
        final_checkpoint["llm_budget"] = payload["llm_budget"]
        _write_json(run_dir / "checkpoint.json", final_checkpoint)
        return payload, run_dir


def load_report(run_dir: Path | str) -> dict[str, Any]:
    return json.loads((Path(run_dir) / "report.json").read_text(encoding="utf-8"))
